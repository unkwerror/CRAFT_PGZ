"""Тесты разбора ТЗ без обращения к LLM: извлечение текста + сборка контекста."""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook

from tender_ingest.documents.extract import UnsupportedDocumentError, extract_text
from tender_ingest.documents.prompt import build_context, build_message


def _docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extract_xlsx() -> None:
    data = _xlsx_bytes([["Наименование", "Значение"], ["Площадь застройки", "4500 кв.м"]])
    r = extract_text("смета.xlsx", None, data)
    assert r.kind == "xlsx"
    assert "лист: Смета" in r.text
    assert "Площадь застройки" in r.text
    assert "4500 кв.м" in r.text
    assert r.low_text is False


def test_extract_docx() -> None:
    data = _docx_bytes(["Площадь объекта 1500 кв.м", "Госэкспертиза за счёт заказчика"])
    r = extract_text("тз.docx", None, data)
    assert r.kind == "docx"
    assert "Площадь объекта 1500" in r.text
    assert "Госэкспертиза" in r.text
    assert r.low_text is False
    assert r.truncated is False


def test_extract_unsupported() -> None:
    with pytest.raises(UnsupportedDocumentError):
        extract_text("тз.txt", "text/plain", b"hello")


def _card() -> SimpleNamespace:
    return SimpleNamespace(
        subject="ПИР школы на 500 мест",
        nmck=None,
        currency=None,
        law="44-ФЗ",
        purchase_method="Открытый конкурс",
        stage="Подача заявок",
        advance_raw="20%",
        securities=None,
        publish_date=None,
        submission_deadline=None,
        region_code="72",
        region_name="Тюменская",
        customer_name="Администрация",
        customer_inn="7202000000",
        delivery_place=None,
        etp=None,
        smp_sono=None,
    )


def test_build_context_includes_card_and_scoring() -> None:
    rel = SimpleNamespace(
        score=75,
        verdict="relevant",
        summary="Профильный ПИР",
        reasoning="ядро профиля",
        red_flags=["срок не указан"],
    )
    ctx = build_context(_card(), rel)
    assert "КАРТОЧКА ЗАКУПКИ" in ctx
    assert "ПИР школы на 500 мест" in ctx
    assert "75/100" in ctx
    assert "срок не указан" in ctx
    msg = build_message(ctx, "текст тз")
    assert "ТЕКСТ ТЗ" in msg
    assert "текст тз" in msg


def test_build_context_without_scoring() -> None:
    ctx = build_context(_card(), None)
    assert "КАРТОЧКА ЗАКУПКИ" in ctx
    assert "БЫСТРЫЙ СКОРИНГ" not in ctx


def test_schemas_have_no_enum_with_type_list() -> None:
    """Structured output отклоняет enum при type ["string","null"] (400 Invalid schema).

    Реальный инцидент 2026-07-08: все разборы ТЗ падали из-за drivers.object_use.
    Nullable-перечисления допустимы только через anyOf.
    """
    from tender_ingest.documents.prompt import BRIEF_SCHEMA, brief_schema_with_card
    from tender_ingest.economics.proposer import PROPOSAL_SCHEMA

    def walk(node: object, path: str) -> None:
        if isinstance(node, dict):
            if "enum" in node and isinstance(node.get("type"), list):
                raise AssertionError(f"enum вместе со списком типов: {path}")
            for key, value in node.items():
                walk(value, f"{path}.{key}")
        elif isinstance(node, list):
            for i, item in enumerate(node):
                walk(item, f"{path}[{i}]")

    for name, schema in [
        ("BRIEF_SCHEMA", BRIEF_SCHEMA),
        ("BRIEF_SCHEMA+card", brief_schema_with_card()),
        ("PROPOSAL_SCHEMA", PROPOSAL_SCHEMA),
    ]:
        walk(schema, name)
